"""Build the Russian research report from verified CSV files and original figures."""
import csv
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'report'
RUN = ROOT / 'guardrails_geometry_results/run_20261018T134157_442709Z'
OLD = ROOT / 'guarails_results/analysis'
HIST = ROOT / 'guardrails_training_history'

def read(p):
    return list(csv.DictReader(p.open()))

MET = {r['metric']:r for r in read(OUT/'data/all_geometry_metrics.csv')}
GROUPS = {(r['model'],r['group']):r for r in read(RUN/'three_way_geometry_by_class.csv')}
SPECTRAL = {(r['model'],r['group']):r for r in read(OUT/'data/spectral_by_class.csv')}
PC = {r['model']:r for r in read(RUN/'pc1_semantics_followup/pc1_summary.csv')}
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

def setup_style(name, size, before=0, after=8, line=1.15, bold=False, color='000000'):
    style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name, style.font.size = 'Arial', Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    pf = style.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    return style

setup_style('Normal',11)
setup_style('Title',26,after=3)
setup_style('Subtitle',11,after=8,color='555555')
setup_style('Heading 1',20,before=20,after=6)
setup_style('Heading 2',16,before=18,after=6)
setup_style('Heading 3',14,before=16,after=4,color='434343')
# Named overrides: dense scientific tables, figure captions, file identifiers.
setup_style('Report Table',9.2,after=0,line=1.06)
setup_style('Report Caption',9, before=4,after=6,line=1.1)
setup_style('Report Source',8.5,before=4,after=4,line=1.05,color='555555')
setup_style('Report Figure',11,after=0,line=1)
setup_style('Report Small',10,after=6,line=1.1)
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next = True

def p(text='', style=None, bold=False):
    para=doc.add_paragraph(style=style)
    para.add_run(text).bold=bold
    return para

def heading(text,level=1):
    return doc.add_heading(text,level)

def page(title=None):
    if title:
        h=heading(title)
        h.paragraph_format.space_before=Pt(0)
        h.paragraph_format.page_break_before=True
    else:
        doc.add_page_break()

def table(headers, rows, widths, source=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.autofit=False
    pr=t._tbl.tblPr
    w=pr.find(qn('w:tblW'));w.set(qn('w:type'),'dxa');w.set(qn('w:w'),str(sum(widths)))
    ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'0');ind.set(qn('w:type'),'dxa');pr.append(ind)
    margins=OxmlElement('w:tblCellMar')
    for edge,value in [('top',80),('bottom',80),('start',120),('end',120)]:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'DADCE0');borders.append(el)
    pr.append(borders)
    for grid,width in zip(t._tbl.tblGrid.gridCol_lst,widths):grid.set(qn('w:w'),str(width))
    for values in [headers]+rows:
        row=t.rows[0] if values is headers else t.add_row()
        for j,(cell,value,width) in enumerate(zip(row.cells,values,widths)):
            cell.width=Inches(width/1440)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(width))
            para=cell.paragraphs[0];para.style=doc.styles['Report Table']
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            para.add_run(str(value)).bold=values is headers
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if values is headers:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    if source:p(source,'Report Source')
    return t

def num(x,kind='f',delta=False):
    x=float(x)
    if kind=='percent':s=f'{100*x:+.2f}' if delta else f'{100*x:.2f}';s+=' п.п.' if delta else '%'
    elif kind=='int':s=f'{x:+.0f}' if delta else f'{x:.0f}'
    elif kind=='rank':s=f'{x:+.2f}' if delta else f'{x:.2f}'
    else:s=f'{x:+.4f}' if delta else f'{x:.4f}'
    return s.replace('.',',').replace('-', '−')

def metric_table(specs):
    rows=[]
    for key,label,kind in specs:
        r=MET[key]
        rows.append([label]+[num(r[m],kind) for m in ['base','ce','finetuned']]+[num(r['joint_minus_base'],kind,True)])
    return table(['Метрика','E5 до','После CE','После\nCE+SupCon','Δ CE+SupCon\nк E5'],rows,[3650,1300,1300,1550,1560])

figures=[]
def figure(path,title,caption,width=6.5):
    path=Path(path)
    idx=len(figures)+1
    cap=p(f'Рисунок {idx}. {title}','Report Caption',bold=True)
    cap.paragraph_format.keep_with_next=True
    pic=p(style='Report Figure')
    pic.paragraph_format.keep_with_next=True
    shape=pic.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set('descr',title+'. '+caption)
    p(caption,'Report Caption')
    figures.append(dict(number=idx,path=str(path.relative_to(ROOT)),title=title))

def statement(title,text):
    heading(title,2)
    p(text)

title=p('Геометрия E5: до и после')
title.paragraph_format.space_after=Pt(3)
for r in title.runs:r.font.size=Pt(26);r.bold=False
p('Исследовательский отчёт • WildGuardTest • 18 октября 2026','Subtitle')
p('Сравнение исходного intfloat/e5-small-v2 с опубликованными CE и CE+SupCon на одних 1 699 запросах. «До» означает предобученный E5 до адаптации к safety-задаче. Основное «после» — CE+SupCon; CE служит отдельным контролем. Это три состояния моделей, а не последовательные этапы одного прогона.','Report Small')
heading('1. Все метрики: сводная таблица',2)
p('Часть A. Глобальная форма пространства и разделение классов. Все векторы L2-нормированы; размерность — 384. Δ — разность CE+SupCon и E5.','Report Small')
global_specs=[
('pc1_share','Доля дисперсии PC1','percent'),
('participation_ratio','Participation ratio, PR','rank'),
('effective_rank','Энтропийный эффективный ранг *','rank'),
('dims_90pct','Компонент для 90% дисперсии *','int'),
('variance_trace','Общий разброс: trace ковариации','f'),
('twonn_id','Локальная размерность TwoNN','rank'),
('mean_cosine_exact','Средний попарный косинус, точный','f'),
('mean_cosine_mc','Средний косинус, 100 000 пар','f'),
('cos_within_safe','Косинус внутри safe','f'),
('cos_within_harm','Косинус внутри harmful','f'),
('cos_safe_harm','Косинус между классами','f'),
('cosine_gap_macro','Косинусный разрыв, macro','f'),
('label_alignment_macro','Внутриклассовая дистанция², macro','f'),
('uniformity_t2','Uniformity, t = 2','f'),
('silhouette_macro','Silhouette по cosine distance, macro *','f'),
('silhouette_micro','Silhouette по cosine distance, micro *','f'),
('within_scatter','Внутриклассовый разброс W *','f'),
('between_scatter','Межклассовый разброс B *','f'),
('between_over_within','Отношение B / W *','f'),
('centroid_cosine','Косинус нормированных центроидов *','f'),
]
metric_table(global_specs)
p('Источники: S1–S3. * Дополнительно пересчитано из сохранённых эмбеддингов. Исходные 13 метрик × 3 модели воспроизведены: максимальное абсолютное расхождение 2,84 × 10⁻¹⁴. PR и энтропийный ранг — разные определения.','Report Source')

page('Сводная таблица — продолжение')
p('Часть B. Локальные окрестности. Purity — доля соседей с той же safety-меткой; macro даёт равный вес двум классам, micro — каждому запросу.','Report Small')
local_specs=[(f'knn_macro@{k}',f'kNN-purity@{k}, macro'+(' *' if k in [1,50,100] else ''),'percent') for k in [1,5,10,20,50,100]]+[
('knn_safe@10','kNN-purity@10: safe','percent'),
('knn_harmful@10','kNN-purity@10: harmful','percent'),
('knn_micro@10','kNN-purity@10, micro *','percent'),
('other_class_closer_fraction','Сосед другого класса ближе своего *','percent'),
('hubness_skew@10','Асимметрия числа вхождений в kNN *','rank')]
metric_table(local_specs)
p('Источники: S2–S3. * Единый пересчёт для всех трёх моделей. Условие «другой класс ближе» использует строгое сравнение расстояний. При равных расстояниях порядок соседей может менять последние знаки.','Report Source')
heading('Как читать сводку',2)
p('Основной результат — рост однородности соседей по safety-метке: 69,27% → 84,33% при k = 10, или +15,05 процентного пункта. Случайный macro-уровень с учётом исключения самого объекта составляет 49,97%.')
p('Одновременно вариация стала значительно сильнее сосредоточена вдоль нескольких направлений: PC1 выросла с 7,07% до 81,76%, PR снизился с 80,62 до 1,50. Эти изменения описывают форму пространства; сами по себе они не означают ни улучшения всей семантики, ни её утраты.')
p('CE и CE+SupCon почти совпадают по purity@10: 84,42% и 84,33%. Разность −0,09 п.п. не даёт основания заявить преимущество SupCon по этой метрике. При этом глобальные косинусы и компактность существенно различаются.')

page('Геометрия внутри каждого класса')
p('Таблица 2. Ковариация центрируется отдельно внутри safe и harmful. Это позволяет отличить разделение двух облаков от изменения формы каждого облака.','Report Small')
class_specs=[('pc1_share','Доля PC1','percent'),('participation_ratio','PR','rank'),('effective_rank','Энтропийный ранг *','rank'),('dims_90pct','Компонент для 90% *','int'),('variance_trace','Trace','f'),('mean_cosine_exact','Средний косинус','f'),('twonn_id','TwoNN','rank'),('uniformity_t2','Uniformity, t=2','f')]
rows=[]
for group,label in [('safe','safe, n = 945'),('harmful','harmful, n = 754')]:
    for key,name,kind in class_specs:
        values=[(SPECTRAL if key in ['effective_rank','dims_90pct'] else GROUPS)[m,group][key] for m in ['base','ce','finetuned']]
        rows.append([f'{label}\n{name}']+[num(x,kind) for x in values])
table(['Класс и метрика','E5 до','После CE','После CE+SupCon'],rows,[4080,1680,1680,1920], 'Источники: S1–S3. Для всех групп доля невалидных наблюдений TwoNN равна 0. * Пересчёт отчёта.')
p('У CE+SupCon относительно E5 trace уменьшается на 76,59% внутри safe и на 42,68% внутри harmful. PR снижается на 94,63% и 97,58% соответственно. У CE trace, напротив, выше исходного в обоих классах. Следовательно, «дообучение сжимает облака» верно здесь для CE+SupCon и не является общим утверждением обо всех objective.','Report Small')

page('2. Основные графики до / после')
p('На исходных рисунках base = E5 до адаптации, finetuned = CE+SupCon. CE показан в сводных таблицах и отдельной диагностике PC1.','Report Small')
figure(RUN/'geometry_metrics_before_after.png','Обзор изменения геометрии','Исходный график S1. Рост purity сопровождается снижением PR и внутриклассовой дистанции. Здесь средний косинус оценён по случайным парам; в таблице также приведено точное значение.')
figure(RUN/'class_geometry_before_after.png','Компактность и форма классов','Исходный график S1. Внутриклассовые косинусы возрастают, а PR и trace уменьшаются. Изменение сильнее для safe по абсолютному сжатию и для harmful по концентрации спектра.')

page('Спектр и устойчивость соседств')
figure(RUN/'cosines_and_spectra.png','Попарные косинусы и накопленная дисперсия','Исходный график S1: сверху распределения косинусов, снизу накопленная доля дисперсии. После CE+SupCon основная вариация сосредоточена в первых компонентах; остаточная многомерная структура сохраняется.')
figure(RUN/'knn_purity_vs_k.png','Эффект сохраняется при изменении k','Исходный график S1. CE+SupCon выше E5 для каждого k от 1 до 100: прирост macro-purity от 8,48 до 23,47 п.п. Пунктир около 0,5 — ожидаемый случайный уровень, не доверительный интервал.')

page('Две PCA-визуализации одних текстов')
figure(RUN/'pca_scatter_before_after.png','Общая плоскость для двух моделей','Исходный график S1. Общая плоскость показывает лишь 1,47% внутримодельной вариации E5 и 81,79% CE+SupCon. Маленькое облако E5 в этой проекции не означает меньший полный разброс. Положение центров между моделями зависит от выбранных координат.')
figure(RUN/'pca_scatter_separate_axes.png','Отдельная PCA для каждой модели','Исходный график S1. Собственные первые две компоненты объясняют 10,40% вариации E5 и 82,74% CE+SupCon. Оси и масштабы двух панелей различаются; метрические выводы делаются по 384D-таблицам.')

page('Что отражает первая компонента')
figure(RUN/'pc1_semantics_followup/pc1_label_and_length.png','Связь PC1 с safety-меткой и длиной','Исходный график S4. PCA рассчитана отдельно для каждой модели. Знак оси развёрнут так, чтобы средняя координата harmful была выше; абсолютный знак PCA не имеет самостоятельного смысла.')
rows=[]
for key,label in [('label_point_biserial_r','Корреляция PC1 с меткой, r'),('label_auc_on_this_sample','AUC метки по координате PC1'),('token_spearman_all','Spearman PC1 с длиной'),('r2_logtokens_only','R²: только log(1+длина)'),('extra_r2_label_given_length','Доп. R² метки при учёте длины'),('extra_r2_length_given_label','Доп. R² длины при учёте метки')]:
    rows.append([label]+[num(PC[m][key]) for m in ['base','ce','finetuned']])
table(['Диагностика PC1','E5','CE','CE+SupCon'],rows,[4320,1560,1680,1800], 'Источник: S4. AUC и R² описывают эту же выборку; это не независимая оценка нового классификатора и не причинные эффекты.')
p('У E5 первая компонента сильно связана с длиной: Spearman ρ = 0,905. После CE+SupCon связь с меткой усиливается до r = 0,758, а метка добавляет 0,543 к R² даже после линейного учёта log-длины. Это поддерживает интерпретацию PC1 как выраженного safety-направления на этой выборке.','Report Small')

page('3. Выводы, которые поддерживают данные')
statement('1. После дообучения соседи стали однороднее по safety-метке.',
'Для CE+SupCon macro kNN-purity@10 выросла с 69,27% до 84,33% (+15,05 п.п.). Улучшение наблюдается на всём проверенном диапазоне k = 1…100. Это означает, что ближайшие тексты чаще принадлежат тому же классу harmful / safe. Purity не является точностью отдельного классификатора и не измеряет тематическую близость. [S1–S3]')
statement('2. Наибольший прирост однородности получен для safe.',
'Purity@10 внутри safe выросла с 63,23% до 87,26% (+24,03 п.п.), внутри harmful — с 75,32% до 81,39% (+6,07 п.п.). Поэтому один macro-показатель скрывает асимметрию: после дообучения harmful-окрестности остаются более смешанными. Доля запросов, для которых чужой класс ближе своего, снизилась с 22,72% до 13,54%. [S2–S3]')
statement('3. Глобальное разделение классов также стало выраженнее.',
'Macro silhouette по косинусному расстоянию выросла с 0,0163 до 0,5251; macro cosine gap — с 0,0033 до 0,1356. Внутри safe и harmful средние косинусы выросли до 0,9520 и 0,8896, тогда как межклассовый косинус изменился мало: 0,7978 → 0,7851. Для CE+SupCon рост разрыва связан главным образом с усилением внутриклассовой близости. [S1–S3]')
statement('4. Пространство стало более анизотропным, но не буквально одномерным.',
'Доля PC1 выросла с 7,07% до 81,76%, PR снизился с 80,62 до 1,50, энтропийный ранг — с 181,51 до 4,05. При этом для 90% дисперсии всё ещё нужны 27 компонент, а TwoNN даёт локальную оценку 15,85. Глобальная концентрация дисперсии и локальная размерность измеряют разные свойства. Низкий PR не равен числу ненулевых координат и не доказывает потери смысла. [S1–S3, M2]')

page('Выводы — продолжение')
statement('5. CE и CE+SupCon дают похожую purity при разной геометрии.',
'При k = 10 macro-purity равна 84,42% для CE и 84,33% для CE+SupCon. У CE гораздо больше cosine gap (0,5536 против 0,1356) и общий trace (0,7464 против 0,1426); у CE+SupCon выше macro silhouette (0,5251 против 0,4648) и сильнее концентрация PC1 (81,76% против 64,55%). Эти метрики отвечают на разные вопросы, поэтому из большей purity или большего разрыва нельзя вывести единое ранжирование качества всех аспектов пространства. [S1–S3]')
statement('6. CE+SupCon сжимает внутриклассовые облака, особенно safe.',
'Внутри safe trace уменьшается до 23,41% от исходного, внутри harmful — до 57,32%. Одновременно PR падает до 5,37% и 2,42% от исходного. Диагностический критерий проекта — и PR, и trace меньше половины значений E5 — срабатывает только для safe. Это выбранный порог для дополнительной проверки; отсутствие флага у harmful не означает отсутствия сжатия. [S1]')
statement('7. Доминирующая ось после обучения тесно связана с safety.',
'У E5 PC1 в основном связана с длиной запросов. Для CE и CE+SupCon связь с safety-меткой намного сильнее и сохраняется в линейном анализе с учётом длины: добавочный R² метки равен 0,532 и 0,543. Это не доказывает, что PC1 кодирует только safety; остальная семантика в этом анализе не размечена. [S4]')
statement('8. Safety-структура сохраняется и вне направления головы.',
'В архивной абляции CE+SupCon удаление направления классификационной головы снижает macro silhouette с 0,5251 до 0,1255 и cosine gap с 0,1356 до 0,0039. Но purity@10 остаётся высокой: 84,33% → 85,30%, при среднем пересечении списков соседей 67,87%. Это свидетельство остаточной структуры, связанной с меткой; улучшение классификатора на остатках не проверялось. [S5]')

page('4. Что именно измеряют метрики')
p('Обозначения: uᵢ — L2-нормированный вектор, C = (1/n)Σ(uᵢ − μ)(uᵢ − μ)ᵀ — центрированная ковариация, λⱼ — её собственные значения, pⱼ = λⱼ / Σλ. Для внутриклассовых строк центр и ковариация вычисляются отдельно. Все выводы о форме относятся к этому пространству.')
statement('Спектр: PC1, PR, энтропийный ранг и n90',
'PC1 = p₁ показывает долю вариации первой компоненты. PR = 1 / Σpⱼ² и эффективный ранг = exp(−Σpⱼ ln pⱼ) описывают концентрацию спектра с разным весом его хвоста. n90 — минимальное число компонент, объясняющих не менее 90% дисперсии. Эти величины нужно читать вместе с общим разбросом trace(C) = Σλⱼ.')
statement('Косинусы, расстояния и разброс',
'Косинус считается между различными объектами; пары объекта с самим собой исключаются. Macro cosine gap — полусумма средних косинусов внутри классов минус средний косинус между классами. Label alignment — среднее квадратов расстояний для пар одной safety-метки, с равным весом классов: 2 − cos_safe − cos_harm. Здесь это alignment по метке, а не по человеческой семантической близости. [M3]')
p('W — взвешенный размерами классов trace внутри классов; B — взвешенный квадрат расстояний их центров до общего центра. Для выбранного деления ковариации на n выполняется trace(C) = W + B. Cosine центроидов использует нормированные средние векторы; он отличается от среднего косинуса всех межклассовых пар.')
statement('Silhouette и локальная однородность',
'Silhouette для объекта — (b − a) / max(a,b), где a — среднее косинусное расстояние до своего класса, b — до другого. Macro усредняет сначала внутри каждого класса. kNN-purity@k — доля совпадающих меток среди k ближайших соседей без самого объекта. На единичной сфере евклидово и косинусное расстояния дают одинаковый порядок соседей.')

page('Методика и границы интерпретации')
statement('Uniformity, TwoNN и hubness',
'Uniformity(t=2) = log mean exp(−2‖uᵢ − uⱼ‖²), оценённая по 100 000 одинаковых случайных пар для всех моделей. Более отрицательные значения означают большую распределённость по этому критерию; изменение само по себе не является показателем качества safety. [M3]')
p('TwoNN оценивает локальную размерность по отношению расстояний до второго и первого соседей. В текущей реализации отброшены верхние 10% log-отношений и выполнена регрессия через начало координат; это оценка, зависящая от данных и масштаба, а не точная размерность пространства. Hubness здесь — асимметрия распределения числа попаданий каждого объекта в чужие списки 10 соседей (scipy skew, bias=True). [M2]')
statement('Единый протокол сравнения',
'После исключения 26 строк WildGuardTest без пригодной метки/текста остаются 945 safe и 754 harmful. Дубликатов нет; три запроса обрезаны до 512 токенов. Для всех моделей: префикс «query: », masked mean последнего слоя со специальными токенами, eval и FP32 при извлечении; float64 и L2-нормировка для геометрии. [S1, M1]')
p('Все 39 исходных значений и дополнительные показатели пересчитаны из сохранённых массивов без нового обучения. Macro даёт равный вес двум классам. [S3]')
statement('Какие выводы пока не установлены',
'Нет повторов обучения по нескольким seed и интервалов для различий между прогонами. Соответствие HF-весов архиву и идентичность всех условий не подтверждены, поэтому причинный эффект только SupCon не выделен. Значимость небольших различий CE и CE+SupCon не установлена.')
p('Шаг и loss возникновения изменений неизвестны: нет промежуточных весов и геометрии фиксированной выборки по шагам. Архивный best step = 1400 обозначает выбор модели и не устанавливает момент сжатия. [S6]')
p('Смысловая близость отдельно не проверена: 0 из 120 подготовленных пар имеют человеческую оценку. При сравнении E5 и CE+SupCon сохраняется в среднем 26,26% соседей@10; это доказывает перестройку соседств, но не потерю смысла. Рост purity следует называть однородностью по safety-метке. [S3–S4]')

page('5. Дополнительная диагностика CE+SupCon')
p('Этот раздел использует архив от 15.10.2026 для той же ревизии CE+SupCon и того же WildGuardTest. Он дополняет единое сравнение трёх моделей. Послойные значения рассчитаны на 512 текстах; остальные приведённые здесь агрегаты — на 1 699. [S5]','Report Small')
figure(OLD/'head_direction_ablation.png','Удаление направления классификационной головы','Слева — доля raw-дисперсии вдоль нормали головы; справа — метрики после проекции raw-векторов на ортогональное подпространство и повторной L2-нормировки. Raw- и unit-показатели не следует смешивать.')
ab=read(OLD/'head_ablation.csv')
rows=[]
for key,label,kind in [('effective_rank','Энтропийный ранг','rank'),('pc1_variance','Доля PC1','percent'),('cosine_gap_macro','Cosine gap, macro','f'),('silhouette_macro','Silhouette, macro','f'),('macro_purity','Purity@10, macro','percent')]:
    rows.append([label,num(ab[0][key],kind),num(ab[1][key],kind)])
table(['Метрика','Исходное пространство','Без направления головы'],rows,[4080,2640,2640], 'Источник: S5, head_ablation.csv. Вдоль головы лежит 80,19% центрированной raw-дисперсии.')
p('Разница между сильным падением глобального разделения и сохранением локальной purity показывает, что разные масштабы структуры ведут себя по-разному. Повторная нормировка после абляции меняет расстояния, поэтому результат не является оценкой готового классификатора на исходных векторах.')

page('Глобальная геометрия и граница головы')
figure(OLD/'geometry_overview.png','Архивный обзор конечной модели','Распределения косинусов, спектр и нормы raw-эмбеддингов до нормировки. В основном сравнении формы используется unit-пространство. График не содержит исходный E5.')
figure(OLD/'head_boundary.png','Локальная purity около границы решения','Ось X — расстояние в raw-пространстве до границы головы. Связь ошибок с перемешанными соседствами носит описательный характер. Порог в архивной диагностике — 0,8917.')
p('При этом пороге AUROC головы = 0,9318, фактический FPR = 2,22%, FNR = 39,79%. Хорошая геометрическая группировка не означает отсутствия пропусков harmful-запросов. На adversarial-подгруппе средняя purity равна 79,22%, на остальных запросах — 89,45%; это micro-средние внутри подгрупп. [S5]','Report Small')

page('Слои и двумерные проекции')
figure(OLD/'layerwise_geometry.png','Геометрия по слоям конечной модели','На одной стратифицированной подвыборке из 512 текстов разделение усиливается в поздних слоях. На графике silhouette_cosine — micro-среднее; его нельзя путать с macro-колонкой основной таблицы. Это разбор слоёв готовой модели, а не история оптимизации.')
figure(OLD/'projections.png','PCA и t-SNE с ошибками головы','Ошибки отмечены крестиками. PCA сохраняет 82,74% дисперсии в двух координатах; trustworthiness@10: PCA = 0,8575, t-SNE = 0,9483. Это качество визуального соседства, а не классификации. Размеры и расстояния островков t-SNE не интерпретируются как глобальные расстояния исходного пространства.')

page('Контроль соседств и история обучения')
figure(OLD/'multiscale_neighborhoods.png','Перестановочный контроль purity','Архив S5: 199 перестановок глобально и внутри adversarial × квартиль длины. При k = 10 условный контроль даёт среднее 0,5200 и 95%-полосу 0,5092–0,5298; наблюдается 0,8433. Полоса — распределение нулевой модели, не доверительный интервал качества и не контроль всех возможных смешивающих факторов.')
figure(HIST/'training_history.png','История архивного CE+SupCon','Архив S6: loss, batch-косинусы, validation-метрики и learning rate. Жирные кривые — скользящие средние 10 логируемых батчей. Вертикаль 1400 отмечает выбранный checkpoint; восстановить по этой картинке динамику PR, TwoNN или purity на фиксированном тесте нельзя.')

page('Контекст: другие objective и OOD')
figure(HIST/'archive_pc1_share.png','PC1 в архиве пяти обучающих режимов','S6: исходный E5, CE, InfoNCE, CE+InfoNCE, SupCon и CE+SupCon. Holdout — WildGuardTest, n = 1 699; геометрия OOD — подвыборка ToxicChat, n = 1 500. Значения округлены в исходном логе; для остальных трёх objective единый пересчёт эмбеддингов в этом отчёте не выполнен.')
figure(HIST/'archive_class_cosines.png','Внутри- и межклассовые косинусы в архиве','S6: качественные различия objective видны также на OOD. Высокая PC1 у SupCon и CE+SupCon сохраняется, но перенос всей геометрии и смыслового качества этим набором показателей не установлен.')
table(['Архивный режим','Лучший шаг','Val AUROC','Val TPR@FPR=1%'],[[r['run'].replace('pure-',''),r['best_step'],num(r['best_val_auroc']),num(r['best_val_tpr'],'percent')] for r in read(HIST/'runs_summary.csv')],[3600,1320,2040,2400], 'S6. Чекпоинты выбраны на разных шагах; все прогоны имеют seed 42. Val-метрики не являются метриками теста.')

page('6. Источники и воспроизводимость')
p('Приоритет доказательств: единый расчёт на сохранённых эмбеддингах → дополнительная диагностика тех же моделей → архивные результаты других objective. Повторяющиеся копии графиков в rq3_history в отчёт второй раз не включены. Все 16 уникальных исходных PNG приведены с подписями.','Report Small')
sources=[
('S1','Основной прогон от 18.10.2026','guardrails_geometry_results/run_20261018T134157_442709Z/','three_way_comparison.csv; three_way_geometry_by_class.csv; manifest.json; data_audit.json; endpoint_diagnostics.csv; pca_2d_metadata.json.'),
('S2','Локальные соседства','Тот же основной прогон','three_way_knn_purity.csv; knn_purity_vs_k.csv. Основная кривая содержит все k = 1…100.'),
('S3','Независимый пересчёт для этого отчёта','report/analyze_report.py; report/data/','all_geometry_metrics.csv; spectral_by_class.csv; reproduction_checks.csv; neighbor_overlap.json; analysis_manifest.json. Последний содержит SHA-256 входных файлов.'),
('S4','PC1 и семантический follow-up','Основной прогон / pc1_semantics_followup/','pc1_summary.csv; PROTOCOL.md; semantic_evaluation.md. Код: rq4_followup/prepare.py и evaluate.py.'),
('S5','Архив подробной диагностики от 15.10.2026','guarails_results/analysis/','geometry.csv; head_ablation.csv; head_axis_variance.csv; head_metrics.csv; layerwise.csv; local_summary.csv; multiscale.csv; projection_quality.csv; group_metrics.csv; metadata.json.'),
('S6','История обучения и другие objective','guardrails_training_history/; rq3_training/','runs_summary.csv; geometry_from_log.csv; lodo_from_archive.csv; *_loss.csv; *_eval.csv; rq3_training/training/TRAINING.md; configs; runs/*/config_used.json и result.json.'),
]
for code,name,path,files in sources:
    p(f'{code}. {name}','Report Small',bold=True)
    p(path+'\n'+files,'Report Source')
p('Основной ноутбук: guardrails_geometry_colab.ipynb. Полный протокол и формулы просмотрены; для отчёта не запускались обучение или повторное извлечение эмбеддингов. Старые и новые data_hash имеют разные схемы формирования, поэтому их различие само по себе не означает различие корпуса.','Report Small')

page('Ревизии моделей и методические ссылки')
manifest=json.loads((RUN/'manifest.json').read_text())
for model,label in [('base','E5 до'),('ce','CE'),('finetuned','CE+SupCon')]:
    entry=manifest['models'][model]
    p(label,'Report Small',bold=True)
    p(entry['cached_repo']+'\ncommit: '+entry['cached_revision'],'Report Source')
p('WildGuardMix / wildguardtest / test','Report Small',bold=True)
p('revision: '+manifest['data']['revision']+'\nОценочный data_hash: '+manifest['data_hash'],'Report Source')
links=[
('M1. intfloat/e5-small-v2 — официальная карточка модели. Префикс query, mean pooling, нормировка, размерность 384.', 'https://huggingface.co/intfloat/e5-small-v2'),
('M2. Facco et al. Estimating the intrinsic dimension of datasets by a minimal neighborhood information. TwoNN.', 'https://arxiv.org/abs/1803.06992'),
('M3. Wang & Isola (2020). Understanding Contrastive Representation Learning through Alignment and Uniformity on the Hypersphere.', 'https://proceedings.mlr.press/v119/wang20k.html'),
('M4. Ai2 WildGuardMix — описание корпуса и разметки. Train использует GPT-4-метки; test — человеческую разметку с дополнительным аудитом.', 'https://huggingface.co/datasets/allenai/wildguardmix')]
for label,url in links:
    para=p(label,'Report Small')
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1155CC');props.append(color);r.append(props)
    txt=OxmlElement('w:t');txt.text=' Открыть источник';r.append(txt);h.append(r);para._p.append(h)
p('Эти данные не позволяют отдельно оценить причинный вклад LLM-источника меток. Для такого вывода нужны одинаковые train-тексты с human- и LLM-разметкой при фиксированных objective и бюджете обучения. [M4, S6]','Report Small')

doc.core_properties.title='Геометрия E5 до и после дообучения'
doc.core_properties.subject='WildGuardTest: E5, CE и CE+SupCon'
doc.core_properties.author=''
doc.core_properties.keywords='E5, geometry, safety, CE, SupCon, kNN, PCA'
target=OUT/'e5_geometry_report.docx'
doc.save(target)
(OUT/'figure_manifest.json').write_text(json.dumps(figures,ensure_ascii=False,indent=2))
(OUT/'design_tokens.json').write_text(json.dumps(dict(preset='google_docs_default',page_inches=[8.5,11],margins_inches=1,table_width_dxa=9360,table_indent_dxa=0,overrides={'Report Table':'Arial 9.2 pt, line 1.06, before/after 0','Report Caption':'Arial 9 pt, before 4 / after 6, line 1.1','Report Source':'Arial 8.5 pt gray, before/after 4, line 1.05','Report Small':'Arial 10 pt, after 6, line 1.1'}),ensure_ascii=False,indent=2))
print(target)
print('Figures:',len(figures),'Tables:',len(doc.tables))
